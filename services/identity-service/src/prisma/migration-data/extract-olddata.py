#!/usr/bin/env python3
"""
One-time extraction: docs/DTI/olddata.xlsx -> oldlist-raw.json

Same pure-cell-dump approach as extract-longlist.py. olddata.xlsx has the
same schema as Longlist.xlsx minus two columns ("Specific Industry" and
"Previous Startup if Applicable") -- those two fields are emitted as null
so the JSON shape matches longlist-raw.json exactly and migrate-longlist.ts's
mapping logic can be reused unchanged.

Run once: .venv/bin/python3 services/identity-service/src/prisma/migration-data/extract-olddata.py
"""
import json
from datetime import datetime, date
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[5]
SRC_XLSX = REPO_ROOT / "docs" / "DTI" / "olddata.xlsx"
OUT_JSON = Path(__file__).resolve().parent / "oldlist-raw.json"

# Column letter -> field key (matches the sheet's header row exactly).
# Same target field names as extract-longlist.py's COLUMNS, minus
# specificIndustry/previousStartup which don't exist in this sheet.
COLUMNS = {
    "A": "startupId",
    "B": "startup",
    "C": "shortDescription",
    "D": "industry",
    "E": "yearFounded",
    "F": "city",
    "G": "province",
    "H": "startupStage",
    "I": "operatingStatus",
    "J": "legalIdentity",
    "K": "revenueStatus",
    "L": "mrr",
    "M": "fundingStage",
    "N": "jobsCreated",
    "O": "marketsReached",
    "P": "primaryContactPerson",
    "Q": "designation",
    "R": "emailAddress",
    "S": "mobileNumber",
    "T": "website",
    "U": "socialMediaLink",
    "V": "programsParticipated",
    "W": "majorMilestones",
    "X": "supportNeeded",
    "Y": "prioritySupportArea",
    "Z": "remarks",
    "AA": "lastUpdated",
}


def cell_to_value(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    return v


def main():
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    ws = wb.active

    rows = []
    for row in ws.iter_rows(min_row=2):
        by_col = {c.column_letter: cell_to_value(c.value) for c in row}
        if not by_col.get("B"):
            continue  # skip placeholder rows with no company name
        record = {field: by_col.get(col) for col, field in COLUMNS.items()}
        record["specificIndustry"] = None
        record["previousStartup"] = None
        rows.append(record)

    print(f"Extracted {len(rows)} rows with a company name (out of {ws.max_row - 1} total data rows)")
    OUT_JSON.write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {OUT_JSON}")


if __name__ == "__main__":
    main()
